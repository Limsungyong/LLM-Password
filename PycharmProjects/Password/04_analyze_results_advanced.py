import csv
import math
import string
from pathlib import Path
from collections import defaultdict, Counter
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


# =========================================================
# 프로젝트 외부 데이터 경로 설정
#
# 예시:
# D:\PasswordData\checked
# D:\PasswordData\complete
# =========================================================

CHECKED_DIR = Path(r"D:\Code\심화과정\논문\data\checked\Desk")
OUTPUT_DIR = Path(r"D:\Code\심화과정\논문\data\complete\Desk")

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

OUTPUT_XLSX = OUTPUT_DIR / "summary_analysis_advanced.xlsx"


def shannon_entropy(password: str) -> float:
    if not password:
        return 0.0

    counts = Counter(password)
    length = len(password)

    entropy = 0.0

    for count in counts.values():
        p = count / length
        entropy -= p * math.log2(p)

    return entropy


def has_upper(password: str) -> bool:
    return any(c.isupper() for c in password)


def has_lower(password: str) -> bool:
    return any(c.islower() for c in password)


def has_digit(password: str) -> bool:
    return any(c.isdigit() for c in password)


def has_special(password: str) -> bool:
    return any(c in string.punctuation for c in password)


def detect_group(csv_file: Path) -> str:
    parts = [p.lower() for p in csv_file.parts]
    name = csv_file.name.lower()

    if "condition" in parts or "condition" in name:
        return "Condition"

    if "simple" in parts or "simple" in name:
        return "Simple"

    return "Unknown"


def analyze_file(csv_file: Path):
    total = 0
    pwned = 0
    total_pwned_count = 0

    password_counter = Counter()

    entropy_sum = 0.0
    upper_count = 0
    lower_count = 0
    digit_count = 0
    special_count = 0

    length_stats = defaultdict(lambda: {
        "total": 0,
        "pwned": 0,
        "entropy_sum": 0.0
    })

    with open(csv_file, "r", encoding="utf-8", errors="ignore") as f:
        reader = csv.DictReader(f)

        for row in reader:
            password = row["password"]
            length = int(row["length"])
            is_pwned = int(row["is_pwned"])
            pwned_count = int(row["pwned_count"])

            total += 1
            password_counter[password] += 1

            entropy = shannon_entropy(password)
            entropy_sum += entropy

            if has_upper(password):
                upper_count += 1

            if has_lower(password):
                lower_count += 1

            if has_digit(password):
                digit_count += 1

            if has_special(password):
                special_count += 1

            length_stats[length]["total"] += 1
            length_stats[length]["entropy_sum"] += entropy

            if is_pwned:
                pwned += 1
                total_pwned_count += pwned_count
                length_stats[length]["pwned"] += 1

    unique_count = len(password_counter)
    duplicate_count = total - unique_count

    return {
        "file_name": csv_file.name,
        "group": detect_group(csv_file),
        "total_passwords": total,
        "unique_passwords": unique_count,
        "duplicate_passwords": duplicate_count,
        "duplicate_ratio_percent": round((duplicate_count / total * 100) if total else 0, 4),
        "unique_ratio_percent": round((unique_count / total * 100) if total else 0, 4),
        "pwned_passwords": pwned,
        "pwned_ratio_percent": round((pwned / total * 100) if total else 0, 4),
        "average_pwned_count": round((total_pwned_count / pwned) if pwned else 0, 2),
        "average_entropy": round((entropy_sum / total) if total else 0, 4),
        "uppercase_ratio_percent": round((upper_count / total * 100) if total else 0, 4),
        "lowercase_ratio_percent": round((lower_count / total * 100) if total else 0, 4),
        "digit_ratio_percent": round((digit_count / total * 100) if total else 0, 4),
        "special_ratio_percent": round((special_count / total * 100) if total else 0, 4),
        "length_stats": length_stats
    }


def style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC")
    )

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter

        for cell in col:
            value = str(cell.value) if cell.value is not None else ""
            max_length = max(max_length, len(value))

        ws.column_dimensions[column].width = min(max_length + 3, 45)


def write_summary_sheet(ws, results):
    headers = [
        "file_name",
        "total_passwords",
        "unique_passwords",
        "duplicate_passwords",
        "duplicate_ratio_percent",
        "unique_ratio_percent",
        "pwned_passwords",
        "pwned_ratio_percent",
        "average_pwned_count",
        "average_entropy",
        "uppercase_ratio_percent",
        "lowercase_ratio_percent",
        "digit_ratio_percent",
        "special_ratio_percent"
    ]

    ws.append(headers)

    for r in results:
        ws.append([
            r["file_name"],
            r["total_passwords"],
            r["unique_passwords"],
            r["duplicate_passwords"],
            r["duplicate_ratio_percent"],
            r["unique_ratio_percent"],
            r["pwned_passwords"],
            r["pwned_ratio_percent"],
            r["average_pwned_count"],
            r["average_entropy"],
            r["uppercase_ratio_percent"],
            r["lowercase_ratio_percent"],
            r["digit_ratio_percent"],
            r["special_ratio_percent"]
        ])

    style_sheet(ws)


def write_length_sheet(ws, results):
    headers = [
        "file_name",
        "length",
        "total_passwords",
        "pwned_passwords",
        "pwned_ratio_percent",
        "average_entropy"
    ]

    ws.append(headers)

    for r in results:
        for length, stats in sorted(r["length_stats"].items()):
            total = stats["total"]
            pwned = stats["pwned"]
            entropy_avg = stats["entropy_sum"] / total if total else 0

            ws.append([
                r["file_name"],
                length,
                total,
                pwned,
                round((pwned / total * 100) if total else 0, 4),
                round(entropy_avg, 4)
            ])

    style_sheet(ws)


def main():
    checked_files = sorted(CHECKED_DIR.rglob("*_checked.csv"))

    if not checked_files:
        print("checked csv 파일이 없습니다.")
        return

    results = [analyze_file(file) for file in checked_files]

    simple_results = [r for r in results if r["group"] == "Simple"]
    condition_results = [r for r in results if r["group"] == "Condition"]
    unknown_results = [r for r in results if r["group"] == "Unknown"]

    wb = Workbook()

    ws_simple = wb.active
    ws_simple.title = "Simple"
    write_summary_sheet(ws_simple, simple_results)

    ws_condition = wb.create_sheet("Condition")
    write_summary_sheet(ws_condition, condition_results)

    ws_simple_length = wb.create_sheet("Simple_Length")
    write_length_sheet(ws_simple_length, simple_results)

    ws_condition_length = wb.create_sheet("Condition_Length")
    write_length_sheet(ws_condition_length, condition_results)

    if unknown_results:
        ws_unknown = wb.create_sheet("Unknown")
        write_summary_sheet(ws_unknown, unknown_results)

        ws_unknown_length = wb.create_sheet("Unknown_Length")
        write_length_sheet(ws_unknown_length, unknown_results)

    wb.save(OUTPUT_XLSX)

    print("=" * 60)
    print("고급 분석 완료")
    print(f"Simple 파일 수: {len(simple_results)}")
    print(f"Condition 파일 수: {len(condition_results)}")
    print(f"Unknown 파일 수: {len(unknown_results)}")
    print(f"저장 위치: {OUTPUT_XLSX}")


if __name__ == "__main__":
    main()